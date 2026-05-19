import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUT_DIR = DATA_DIR / "xlsx"

FILES = [
    "user_input_samples.csv",
    "market_transactions_mvp.csv",
    "building_registry_mvp.csv",
    "seed_market.csv",
    "seed_building_registry.csv",
]

NUMERIC_FIELDS = {
    "deposit_won",
    "monthly_rent_won",
    "price_won",
    "price_or_deposit_won",
    "exclusive_area_m2",
    "floor",
    "built_year",
    "approval_year",
}

TEXT_FIELDS = {
    "legal_dong_code",
    "address",
    "road_address",
    "region",
    "contract_stage",
}


def coerce_value(field, value):
    if field in TEXT_FIELDS:
        return value
    if field in NUMERIC_FIELDS and value not in ("", None):
        try:
            if "." in value:
                return float(value)
            return int(value)
        except ValueError:
            return value
    if value in ("true", "false"):
        return value.upper()
    return value


def export_one(path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        headers = rows[0].keys() if rows else next(csv.reader([path.read_text(encoding="utf-8-sig").splitlines()[0]]))

    wb = Workbook()
    ws = wb.active
    ws.title = path.stem[:31]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")

    for row_index, row in enumerate(rows, 2):
        for col_index, header in enumerate(headers, 1):
            value = coerce_value(header, row.get(header, ""))
            cell = ws.cell(row=row_index, column=col_index, value=value)
            if header in TEXT_FIELDS:
                cell.number_format = "@"
            elif header in NUMERIC_FIELDS:
                cell.number_format = "#,##0.###"

    ws.freeze_panes = "A2"
    for col_index, header in enumerate(headers, 1):
        values = [str(header)] + [str(row.get(header, "")) for row in rows]
        width = min(max(len(value) for value in values) + 2, 42)
        ws.column_dimensions[get_column_letter(col_index)].width = width

    OUT_DIR.mkdir(exist_ok=True)
    output = OUT_DIR / f"{path.stem}.xlsx"
    wb.save(output)
    return output


def main():
    for name in FILES:
        output = export_one(DATA_DIR / name)
        print(output.relative_to(ROOT))


if __name__ == "__main__":
    main()
